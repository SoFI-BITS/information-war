"""Build Phase 1 deliverables (Markdown + XLSX) from app/src/scenario/scenario.json.

Usage: python -I tools/build_phase1.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SCN = json.loads((ROOT / "app" / "src" / "scenario" / "scenario.json").read_text(encoding="utf-8"))
OUT = ROOT / "docs" / "phase-1-content"
OUT.mkdir(parents=True, exist_ok=True)

P = SCN["params"]
STOCKS = SCN["stocks"]
ROUNDS = SCN["rounds"]
GROUPS = P["groups"]
SOURCES = SCN["sources"]
TICKERS = [s["ticker"] for s in STOCKS]


def fundamental_path():
    path = {}
    for s in STOCKS:
        f = [float(s["startPrice"])]
        for r in ROUNDS:
            f.append(f[-1] * (1 + r["fundamentals"][s["ticker"]] / 100))
        path[s["ticker"]] = f
    return path


FP = fundamental_path()


def all_signals():
    for r in ROUNDS:
        for sig in r["signals"]:
            yield r, sig


def actual_move(sig):
    return ROUNDS[sig["targetRound"] - 1]["fundamentals"][sig["ticker"]]


def dir_word(d):
    return "Up" if d > 0 else "Down"


# ---------------------------------------------------------------- Markdown
def build_markdown():
    lines = ["# Phase 1 — The Ten Stocks", "",
             f"Setting: **{SCN['meta']['setting']}** (fictional). Currency: **{SCN['meta']['currency']}**.", "",
             f"Every team starts with **{P['endowmentPerStock']:,} credits of each stock** (shares rounded) plus cash "
             f"(about {P['startingCash']:,}) that tops the total up to exactly **{P['startingCapital']:,}**. Position cap: "
             f"**{P['maxPositionValuePerStock']:,} credits per stock at its start price** (fixed share cap below).", "",
             "| # | Ticker | Company | Sector | Start price | Volatility | Starting shares | Max shares |",
             "|---|---|---|---|---:|---|---:|---:|"]
    for i, s in enumerate(STOCKS, 1):
        endow = round(P["endowmentPerStock"] / s["startPrice"])
        cap = int(P["maxPositionValuePerStock"] // s["startPrice"])
        lines.append(f"| {i} | **{s['ticker']}** | {s['name']} | {s['sector']} | {s['startPrice']} | {s['volatility']} | {endow} | {cap} |")
    lines += ["", "## Company profiles (participant-facing)", ""]
    for s in STOCKS:
        lines += [f"### {s['ticker']} — {s['name']}", f"*{s['sector']} · start price {s['startPrice']} · volatility: {s['volatility']}*", "",
                  s["description"], "", f"**Known links:** {s['links']}", ""]
    lines += ["## Volatility disclosure", "",
              "All ten stocks are **high-volatility**. Typical round-to-round reference-value moves are within ±10%; "
              "'Very high' stocks can move ±20% on a major event. This is disclosed to participants; the direction and timing of moves are not.", ""]
    (OUT / "stocks.md").write_text("\n".join(lines), encoding="utf-8")

    m = ["# Phase 1 — Ten-Round Narrative Matrix (ORGANISER ONLY — CONFIDENTIAL)", "",
         "> This file reveals the hidden economic state and the truth of every private signal. "
         "Do not share with participants before the debrief.", "",
         "## How private intelligence is allocated", "",
         f"Teams are placed in **{len(GROUPS)} intel groups ({', '.join(GROUPS)})** by team number: "
         "team 1 → A, team 2 → B, … team 6 → F, team 7 → A, and so on. With 20–30 teams each group has 3–5 teams.", "",
         "Every round has three private signals. Each signal goes to **two groups** (about a third of all teams), so every team "
         "receives exactly one signal per round, but no two groups see the same set of signals across the event. "
         "Pairings were chosen by `tools/balance_signals.py` to balance each group's expected value.", "",
         "Reliability classes: **ACCURATE** · **INCOMPLETE** (true but missing the key fact) · **OUTDATED** (was true; no longer) · **MISLEADING** (false or wrong meaning).", "",
         "## Hidden fundamental path (reference-value anchor)", "",
         "Percent change of each company's hidden fundamental value per round. The reference value tracks this, "
         f"plus a capped {int(P['referenceValue']['weight']*100)}% pull toward the round's traded VWAP (max ±{P['referenceValue']['capPct']}%).", "",
         "| Ticker | Start | " + " | ".join(f"R{r['round']}" for r in ROUNDS) + " | Final F | Total |",
         "|---|---:|" + "---:|" * len(ROUNDS) + "---:|---:|"]
    for s in STOCKS:
        t = s["ticker"]
        f = FP[t]
        cells = " | ".join(f"{r['fundamentals'][t]:+d}%" for r in ROUNDS)
        m.append(f"| {t} | {s['startPrice']} | {cells} | {f[-1]:.2f} | {(f[-1]/f[0]-1)*100:+.1f}% |")
    m += [""]

    m += ["## Round-by-round matrix", ""]
    for r in ROUNDS:
        m += [f"### Round {r['round']} — {r['title']}", "", f"**Purpose:** {r['purpose']}", "",
              f"**Hidden state:** {r['hiddenState']}", "", "**Common news (all teams):**", ""]
        for n in r["news"]:
            m.append(f"- **{n['headline']}** — {n['body']}")
        moves = ", ".join(f"{t} {r['fundamentals'][t]:+d}%" for t in TICKERS)
        m += ["", f"**Fundamental moves:** {moves}", ""]
        if r.get("marketMaker"):
            mm = r["marketMaker"]
            m += [f"**Liquidity override:** market-maker half-spread {mm['halfSpreadPct']}%, {mm['levels']} levels × "
                  f"{mm['levelNotional']:,} credits, skew {mm['skewPctPer10k']}% per 10k net flow (cap {mm['maxSkewPct']}%).", ""]
        if r.get("forecast"):
            m += [f"**Forecast question:** {r['forecast']['question']}", ""]
        m += ["| Signal | Groups | Ticker | Source (shown) | Text shown to teams | Implies | True class | Actual move | Truth (debrief) |",
              "|---|---|---|---|---|---|---|---:|---|"]
        for sig in r["signals"]:
            tgt = "" if sig["targetRound"] == r["round"] else f" (R{sig['targetRound']})"
            m.append(f"| {sig['id']} | {', '.join(sig['groups'])} | {sig['ticker']} | {SOURCES[sig['source']]['label']} | "
                     f"{sig['text']} | {dir_word(sig['impliedDirection'])}{tgt} | {sig['reliability']} | {actual_move(sig):+d}% | {sig['truth']} |")
        m += ["", "**Research available (1 credit each):**", ""]
        for res in r["research"]:
            m.append(f"- `{res['id']}` *{res['title']}* → {res['result']}")
        m.append("")

    m += ["## Allocation matrix (signal received by each group)", "",
          "| Round | " + " | ".join(GROUPS) + " |", "|---|" + "---|" * len(GROUPS)]
    for r in ROUNDS:
        row = []
        for g in GROUPS:
            sig = next(s for s in r["signals"] if g in s["groups"])
            row.append(f"{sig['id']} ({sig['reliability'][:3]})")
        m.append(f"| {r['round']} | " + " | ".join(row) + " |")
    m += ["", "## Balance check", "",
          "Edge = implied direction × actual fundamental move (percentage points). A team that traded every signal "
          "naively at equal size would earn roughly proportional to this.", "",
          "| Group | Signals | Accurate | Incomplete | Outdated | Misleading | Total edge |", "|---|---:|---:|---:|---:|---:|---:|"]
    for g in GROUPS:
        sigs = [s for _, s in all_signals() if g in s["groups"]]
        cnt = {k: sum(s["reliability"] == k for s in sigs) for k in ("ACCURATE", "INCOMPLETE", "OUTDATED", "MISLEADING")}
        edge = sum(s["impliedDirection"] * actual_move(s) for s in sigs)
        m.append(f"| {g} | {len(sigs)} | {cnt['ACCURATE']} | {cnt['INCOMPLETE']} | {cnt['OUTDATED']} | {cnt['MISLEADING']} | {edge:+d} |")
    m += ["", "Stated source reliability vs. actual hit rate (shown to teams as a label on every signal):", "",
          "| Source label | Stated | Signals | Accurate | Hit rate |", "|---|---|---:|---:|---:|"]
    for key, src in SOURCES.items():
        sigs = [s for _, s in all_signals() if s["source"] == key]
        acc = sum(s["reliability"] == "ACCURATE" for s in sigs)
        m.append(f"| {src['label']} | {src['statedReliability']} | {len(sigs)} | {acc} | {acc/len(sigs)*100:.0f}% |")
    m.append("")
    (OUT / "narrative-matrix.md").write_text("\n".join(m), encoding="utf-8")


# ---------------------------------------------------------------- XLSX
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill, cell.alignment = BOLD, HEAD_FILL, Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"


def put(ws, row, col, value, font=BLACK, fmt=None, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = WRAP
    return c


def build_xlsx():
    wb = Workbook()

    # README
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 110
    readme = [
        "THE INFORMATION WAR — Phase 1 scenario workbook (ORGANISER ONLY)",
        "",
        "Generated from app/src/scenario/scenario.json by tools/build_phase1.py. The JSON file is the source the server loads;",
        "if you change numbers here, copy them back into scenario.json (or edit the JSON and regenerate).",
        "",
        "Legend: blue text = input you may edit · black = formula · green = formula pulling from another sheet.",
        "Percent inputs are stored as fractions (0.05 = 5%).",
        "",
        "Sheets: Parameters · Stocks · Hidden State · Public News · Private Signals · Allocation · Balance Check · Research · Forecasts",
        "Intel groups: team 1 → A, team 2 → B, … team 7 → A again (round-robin by team number).",
    ]
    for i, line in enumerate(readme, 1):
        put(ws, i, 1, line, Font(name=FONT, bold=(i == 1), size=12 if i == 1 else 10))

    # Parameters
    ps = wb.create_sheet("Parameters")
    header(ps, ["Parameter", "Value", "Notes"], [34, 14, 70])
    rv, mm = P["referenceValue"], P["marketMaker"]
    params = [
        ("Starting capital", P["startingCapital"], "Total value each team starts with (cash + share endowment).", "#,##0"),
        ("Starting cash", P["startingCash"], "Nominal cash component. Actual cash = capital − endowment value, so every team starts at exactly the starting capital.", "#,##0"),
        ("Endowment per stock", P["endowmentPerStock"], "Credits of each stock given at start (shares = value / start price, rounded).", "#,##0"),
        ("Max position value per stock", P["maxPositionValuePerStock"], "Share cap = INT(this / start price). Fixed for the event.", "#,##0"),
        ("Max order notional", P["maxOrderNotional"], "Quantity × price per order.", "#,##0"),
        ("Max open orders per team", P["maxOpenOrders"], "", "0"),
        ("Market order collar", P["marketOrderCollarPct"] / 100, "Market orders fill only within this % of the best opposite price at submission.", "0.0%"),
        ("Research credits per team", P["researchCredits"], "For the whole event; separate from cash.", "0"),
        ("Rounds", P["rounds"], "", "0"),
        ("News phase (s)", P["newsSeconds"], "", "0"),
        ("Trading phase (s)", P["tradingSeconds"], "", "0"),
        ("Settlement phase (s)", P["settleSeconds"], "", "0"),
        ("RV weight on VWAP gap", rv["weight"], "RV = F + weight × clamp(VWAP − F, ±cap × F).", "0.00"),
        ("RV cap", rv["capPct"] / 100, "", "0.0%"),
        ("RV min volume (shares)", rv["minVolumeShares"], "VWAP ignored below this round volume.", "0"),
        ("RV min trades", rv["minTrades"], "", "0"),
        ("RV min distinct teams", rv["minDistinctTeams"], "", "0"),
        ("MM half-spread", mm["halfSpreadPct"] / 100, "Market-maker best bid/ask distance from its mid.", "0.0%"),
        ("MM level step", mm["levelStepPct"] / 100, "", "0.0%"),
        ("MM levels per side", mm["levels"], "", "0"),
        ("MM level notional", mm["levelNotional"], "", "#,##0"),
        ("MM skew per 10k net flow", mm["skewPctPer10k"] / 100, "Mid shifts this much per 10,000 credits teams net-buy from the MM.", "0.0%"),
        ("MM max skew", mm["maxSkewPct"] / 100, "", "0.0%"),
    ]
    pref = {}
    for i, (name, val, note, fmt) in enumerate(params, 2):
        put(ps, i, 1, name)
        put(ps, i, 2, val, BLUE, fmt)
        put(ps, i, 3, note, wrap=True)
        pref[name] = f"Parameters!$B${i}"

    # Stocks
    ss = wb.create_sheet("Stocks")
    header(ss, ["Ticker", "Company", "Sector", "Description", "Start price", "Volatility", "Known links", "Starting shares", "Max shares", "Start holding value"],
           [8, 26, 24, 60, 11, 11, 50, 12, 11, 14])
    for i, s in enumerate(STOCKS, 2):
        put(ss, i, 1, s["ticker"], Font(name=FONT, bold=True))
        put(ss, i, 2, s["name"])
        put(ss, i, 3, s["sector"])
        put(ss, i, 4, s["description"], wrap=True)
        put(ss, i, 5, s["startPrice"], BLUE, "0.00")
        put(ss, i, 6, s["volatility"], BLUE)
        put(ss, i, 7, s["links"], wrap=True)
        put(ss, i, 8, f"=ROUND({pref['Endowment per stock']}/E{i},0)", GREEN, "0")
        put(ss, i, 9, f"=INT({pref['Max position value per stock']}/E{i})", GREEN, "0")
        put(ss, i, 10, f"=H{i}*E{i}", BLACK, "#,##0")
    n = len(STOCKS) + 1
    put(ss, n + 1, 9, "Total", Font(name=FONT, bold=True))
    put(ss, n + 1, 10, f"=SUM(J2:J{n})", BLACK, "#,##0")
    put(ss, n + 2, 9, "+ cash", Font(name=FONT, bold=True))
    put(ss, n + 2, 10, f"={pref['Starting capital']}-J{n+1}", GREEN, "#,##0")
    ss.cell(row=n + 2, column=10).comment = Comment("Cash tops the endowment up to exactly the starting capital (≈ starting cash after share rounding).", "Organiser")
    put(ss, n + 3, 9, "Start equity", Font(name=FONT, bold=True))
    put(ss, n + 3, 10, f"=J{n+1}+J{n+2}", BLACK, "#,##0")

    # Hidden State
    hs = wb.create_sheet("Hidden State")
    nr = len(ROUNDS)
    cols = ["Ticker", "Start"] + [f"R{r['round']} %" for r in ROUNDS] + ["F0"] + [f"F{r['round']}" for r in ROUNDS] + ["Total return"]
    header(hs, cols, [8, 9] + [7] * nr + [9] * (nr + 1) + [12])
    for i, s in enumerate(STOCKS, 2):
        put(hs, i, 1, s["ticker"], Font(name=FONT, bold=True))
        put(hs, i, 2, f"=Stocks!E{i}", GREEN, "0.00")
        for j, r in enumerate(ROUNDS):
            put(hs, i, 3 + j, r["fundamentals"][s["ticker"]] / 100, BLUE, "0%;-0%;-")
        f0_col = 3 + nr
        put(hs, i, f0_col, f"=B{i}", BLACK, "0.00")
        for j in range(nr):
            col = f0_col + 1 + j
            prev = get_column_letter(col - 1)
            pct = get_column_letter(3 + j)
            put(hs, i, col, f"={prev}{i}*(1+{pct}{i})", BLACK, "0.00")
        last = get_column_letter(f0_col + nr)
        put(hs, i, f0_col + nr + 1, f"={last}{i}/B{i}-1", BLACK, "0.0%")
    hs.cell(row=1, column=3).comment = Comment("Hidden fundamental change per round (organiser input). Source: scenario.json.", "Organiser")

    # Public news
    ns = wb.create_sheet("Public News")
    header(ns, ["Round", "Round title", "Headline", "Body", "Purpose", "Hidden state (organiser only)"], [7, 22, 50, 70, 50, 60])
    row = 2
    for r in ROUNDS:
        for n_ in r["news"]:
            put(ns, row, 1, r["round"])
            put(ns, row, 2, r["title"])
            put(ns, row, 3, n_["headline"], BLUE, wrap=True)
            put(ns, row, 4, n_["body"], BLUE, wrap=True)
            put(ns, row, 5, r["purpose"], wrap=True)
            put(ns, row, 6, r["hiddenState"], wrap=True)
            row += 1

    # Private signals
    sg = wb.create_sheet("Private Signals")
    header(sg, ["Signal", "Round", "Groups", "Ticker", "Source (shown)", "Stated reliability", "Text shown to teams", "Implied dir (+1/-1)",
                "Target round", "True class", "Actual move", "Edge (pts)", "Truth (debrief)"],
           [8, 7, 8, 8, 16, 12, 60, 10, 8, 13, 10, 10, 60])
    row = 2
    for r, sig in all_signals():
        put(sg, row, 1, sig["id"], Font(name=FONT, bold=True))
        put(sg, row, 2, r["round"])
        put(sg, row, 3, ", ".join(sig["groups"]), BLUE)
        put(sg, row, 4, sig["ticker"], BLUE)
        put(sg, row, 5, SOURCES[sig["source"]]["label"], BLUE)
        put(sg, row, 6, SOURCES[sig["source"]]["statedReliability"])
        put(sg, row, 7, sig["text"], BLUE, wrap=True)
        put(sg, row, 8, sig["impliedDirection"], BLUE, "+0;-0")
        put(sg, row, 9, sig["targetRound"], BLUE)
        put(sg, row, 10, sig["reliability"], BLUE)
        # Actual move: look up the ticker row and the target round's % column in Hidden State.
        put(sg, row, 11, f"=INDEX('Hidden State'!$C$2:${get_column_letter(2+nr)}${len(STOCKS)+1},"
                         f"MATCH(D{row},'Hidden State'!$A$2:$A${len(STOCKS)+1},0),I{row})", GREEN, "0%;-0%;-")
        put(sg, row, 12, f"=H{row}*K{row}*100", BLACK, "+0;-0;0")
        put(sg, row, 13, sig["truth"], wrap=True)
        row += 1
    last_sig = row - 1

    # Allocation
    al = wb.create_sheet("Allocation")
    header(al, ["Round"] + GROUPS, [8] + [14] * len(GROUPS))
    for i, r in enumerate(ROUNDS, 2):
        put(al, i, 1, r["round"])
        for j, g in enumerate(GROUPS):
            sig = next(s for s in r["signals"] if g in s["groups"])
            put(al, i, 2 + j, f"{sig['id']} · {sig['reliability'][:3]}")
    put(al, len(ROUNDS) + 3, 1, "Team → group: team number n goes to group index ((n − 1) mod 6). E.g. teams 1, 7, 13, 19, 25 → A.")

    # Balance check
    bc = wb.create_sheet("Balance Check")
    header(bc, ["Group", "Signals", "Accurate", "Incomplete", "Outdated", "Misleading", "Total edge (pts)"], [8, 9, 10, 11, 10, 11, 15])
    rng_g = f"'Private Signals'!$C$2:$C${last_sig}"
    rng_r = f"'Private Signals'!$J$2:$J${last_sig}"
    rng_e = f"'Private Signals'!$L$2:$L${last_sig}"
    for i, g in enumerate(GROUPS, 2):
        put(bc, i, 1, g, Font(name=FONT, bold=True))
        put(bc, i, 2, f'=COUNTIF({rng_g},"*{g}*")', GREEN)
        for j, cls in enumerate(["ACCURATE", "INCOMPLETE", "OUTDATED", "MISLEADING"]):
            put(bc, i, 3 + j, f'=COUNTIFS({rng_g},"*{g}*",{rng_r},"{cls}")', GREEN)
        put(bc, i, 7, f'=SUMIF({rng_g},"*{g}*",{rng_e})', GREEN, "+0;-0;0")
    lr = len(GROUPS) + 1
    put(bc, lr + 2, 1, "Edge spread (max − min)")
    put(bc, lr + 2, 7, f"=MAX(G2:G{lr})-MIN(G2:G{lr})", BLACK, "0")
    put(bc, lr + 3, 1, "Target: spread ≤ 6 pts and accurate counts within 1 of each other.")

    # Research
    rs = wb.create_sheet("Research")
    header(rs, ["Research ID", "Round", "Title", "Cost (credits)", "Result revealed"], [11, 7, 40, 12, 90])
    row = 2
    for r in ROUNDS:
        for res in r["research"]:
            put(rs, row, 1, res["id"])
            put(rs, row, 2, r["round"])
            put(rs, row, 3, res["title"], BLUE)
            put(rs, row, 4, 1, BLUE)
            put(rs, row, 5, res["result"], BLUE, wrap=True)
            row += 1

    # Forecasts
    fc = wb.create_sheet("Forecasts")
    header(fc, ["Round", "Ticker", "Question", "Correct answer", "Points"], [7, 8, 80, 14, 8])
    row = 2
    for r in ROUNDS:
        if r.get("forecast"):
            t = r["forecast"]["ticker"]
            put(fc, row, 1, r["round"])
            put(fc, row, 2, t, BLUE)
            put(fc, row, 3, r["forecast"]["question"], BLUE, wrap=True)
            put(fc, row, 4, "Up (fundamental)" if r["fundamentals"][t] > 0 else "Down (fundamental)")
            put(fc, row, 5, P["forecastPoints"], BLUE)
            row += 1
    put(fc, row + 1, 1, "Scoring is on the published reference value (RV end vs RV start), which includes the small VWAP pull; the fundamental column is the expected answer.")

    wb.save(OUT / "information-war-phase1.xlsx")


if __name__ == "__main__":
    build_markdown()
    build_xlsx()
    print("wrote", OUT)
